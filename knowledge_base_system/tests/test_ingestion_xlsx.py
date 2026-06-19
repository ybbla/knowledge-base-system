import io
import time

from openpyxl import Workbook

from app.core.models import (
    Asset,
    AssetType,
    Document,
    ElementType,
    KnowledgeChunk,
    ParsedElement,
)
from ingestion.pipeline import IngestionPipeline
import ingestion.pipeline as ingestion_module
from parsers.base import DocumentParser, ParseResult
from parsers.registry import ParserRegistry
from parsers.xlsx_parser import XlsxParser


def _xlsx_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "状态表"
    ws.append(["状态", "说明"])
    ws.append(["处理中", "系统正在解析文档"])
    ws.append(["成功", "文档已进入知识库"])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


class _FakeExtractor:
    def __init__(self) -> None:
        self.seen_elements = []

    def extract(self, elements, assets, ingest_job_id, category):
        self.seen_elements = list(elements)
        source = next(
            (el for el in elements if el.element_type == ElementType.table),
            elements[0],
        )
        return [
            KnowledgeChunk(
                doc_id=source.doc_id,
                doc_version=source.doc_version,
                title="XLSX 状态表",
                content="XLSX 表格说明文档解析状态，包括处理中和成功。",
                category=category,
                ingest_job_id=ingest_job_id,
            )
        ]


class _RecordingIndex:
    def __init__(self) -> None:
        self.batches = []

    def add_batch(self, items):
        self.batches.append(items)

    def add(self, *_args, **_kwargs):
        raise AssertionError("入库管线应使用批量索引写入")

    def delete(self, _chunk_id):
        return None

    def search(self, *_args, **_kwargs):
        return []


class _RecordingChunkStore:
    def __init__(self) -> None:
        self.chunks = {}
        self.index_statuses = {}

    def put(self, chunk):
        self.chunks[chunk.chunk_id] = chunk

    def update_index_status(self, chunk_ids, status, error=None):
        for chunk_id in chunk_ids:
            self.index_statuses[chunk_id] = status


class _RecordingAssetStore:
    def __init__(self) -> None:
        self.assets = {}

    def put(self, asset):
        self.assets[asset.asset_id] = asset

    def get(self, asset_id):
        return self.assets.get(asset_id)

    def delete(self, asset_id):
        self.assets.pop(asset_id, None)


class _FakeEmbedder:
    def __init__(self) -> None:
        self.calls = []

    def embed_text(self, texts):
        self.calls.append(list(texts))
        return [[1.0, 0.0] for _ in texts]


class _EmbeddedParser(DocumentParser):
    SUPPORTED_TYPES = {"embedded-test"}

    def supports(self, source_type: str) -> bool:
        return source_type == "embedded-test"

    def parse(self, doc):
        if doc.parent_doc_id:
            elements = [
                ParsedElement(
                    doc_id=doc.doc_id,
                    doc_version=doc.version,
                    sequence_order=1,
                    element_type=ElementType.paragraph,
                    text="子文档内容",
                )
            ]
            doc.source_hash = "sha256:child"
        else:
            asset = Asset(
                doc_id=doc.doc_id,
                asset_type=AssetType.video,
                original_uri="https://example.com/root.mp4",
            )
            elements = [
                ParsedElement(
                    doc_id=doc.doc_id,
                    doc_version=doc.version,
                    sequence_order=1,
                    element_type=ElementType.title,
                    text="根文档",
                ),
                ParsedElement(
                    doc_id=doc.doc_id,
                    doc_version=doc.version,
                    sequence_order=2,
                    element_type=ElementType.paragraph,
                    text="嵌入文档入口",
                    embedded_doc_id="doc_child",
                ),
            ]
            doc.source_hash = "sha256:root"
            return ParseResult(doc=doc, elements=elements, assets=[asset])
        return ParseResult(doc=doc, elements=elements)


def test_ingestion_pipeline_dispatches_xlsx_parser(monkeypatch):
    embedder = _FakeEmbedder()
    monkeypatch.setattr(ingestion_module, "embedding_client", embedder)

    registry = ParserRegistry()
    registry.register(XlsxParser())
    extractor = _FakeExtractor()
    vector_index = _RecordingIndex()
    bm25_index = _RecordingIndex()
    chunk_store = _RecordingChunkStore()
    asset_store = _RecordingAssetStore()
    pipeline = IngestionPipeline(
        parser_registry=registry,
        extractor=extractor,
        vector_index=vector_index,
        bm25_index=bm25_index,
        asset_store=asset_store,
        chunk_store=chunk_store,
    )
    doc = Document(
        title="状态表",
        source_type="xlsx",
        source_uri="memory://status.xlsx",
        category="验收",
        metadata={"raw_content": _xlsx_bytes()},
    )
    doc.ingest_job_id = doc.doc_id

    job = pipeline.submit(doc)
    for _ in range(50):
        if job.status in {"completed", "failed"}:
            break
        time.sleep(0.02)

    assert job.status == "completed", job.error
    assert job.doc_ids == [doc.doc_id]
    assert job.chunk_count == 1
    assert [el.element_type for el in extractor.seen_elements] == [
        ElementType.title,
        ElementType.table,
    ]
    assert len(chunk_store.chunks) == 1
    assert len(embedder.calls) == 1
    assert len(vector_index.batches) == 1
    assert len(bm25_index.batches) == 1


def test_ingestion_pipeline_recurses_embedded_docs_without_repeating_root(monkeypatch):
    embedder = _FakeEmbedder()
    monkeypatch.setattr(ingestion_module, "embedding_client", embedder)

    registry = ParserRegistry()
    registry.register(_EmbeddedParser())
    extractor = _FakeExtractor()
    vector_index = _RecordingIndex()
    bm25_index = _RecordingIndex()
    chunk_store = _RecordingChunkStore()
    asset_store = _RecordingAssetStore()
    pipeline = IngestionPipeline(
        parser_registry=registry,
        extractor=extractor,
        vector_index=vector_index,
        bm25_index=bm25_index,
        asset_store=asset_store,
        chunk_store=chunk_store,
    )
    doc = Document(
        title="Root",
        source_type="embedded-test",
        source_uri="memory://root",
        metadata={"raw_content": "root"},
    )
    doc.ingest_job_id = doc.doc_id

    job = pipeline.submit(doc)
    for _ in range(50):
        if job.status in {"completed", "failed"}:
            break
        time.sleep(0.02)

    assert job.status == "completed", job.error
    assert job.doc_ids == [doc.doc_id, "doc_child"]
    assert job.asset_count == 1
    assert [el.text for el in extractor.seen_elements] == [
        "根文档",
        "嵌入文档入口",
        "子文档内容",
    ]
    assert [asset.original_uri for asset in asset_store.assets.values()] == [
        "https://example.com/root.mp4"
    ]
